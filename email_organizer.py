import os
import base64
import re
import json
import requests
import time

from datetime import datetime, timezone
from email.utils import parsedate_to_datetime
from html import unescape

from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


# ============================================================
# SRNR GMAIL AI ORGANIZER - VERSION 2
# ============================================================
#
# PURPOSE
# -------
# 1. Read Gmail messages.
# 2. Classify each message using LOCAL Ollama/Qwen.
# 3. Apply one controlled Gmail category label.
# 4. Keep emails up to 30 days old in Inbox.
# 5. Archive emails older than 30 days by removing INBOX label.
# 6. Never delete emails.
#
# IMPORTANT
# ---------
# This script requires gmail.modify permission because it applies labels
# and removes the INBOX label from older messages.
# ============================================================


# ============================================================
# CONFIGURATION
# ============================================================

PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))

CREDENTIALS_FILE = os.path.join(
    PROJECT_DIR,
    "credentials",
    "credentials.json"
)

TOKEN_FILE = os.path.join(
    PROJECT_DIR,
    "credentials",
    "token.json"
)

REPORT_DIR = os.path.join(
    PROJECT_DIR,
    "reports"
)

os.makedirs(REPORT_DIR, exist_ok=True)

OLLAMA_URL = "http://127.0.0.1:11434/api/generate"
OLLAMA_MODEL = "qwen2.5:3b"

# Gmail parent label.
# Child labels will look like:
# AI Mail/Business
# AI Mail/Banking
# AI Mail/Personal
LABEL_ROOT = "AI Mail"

# Inbox retention rule.
INBOX_RETENTION_DAYS = 30

# ------------------------------------------------------------
# SAFETY SWITCH
# ------------------------------------------------------------
#
# True  = classify and SHOW what would happen, but do not change Gmail.
# False = actually create/apply labels and archive >30-day Inbox mail.
#
DRY_RUN = True

# ------------------------------------------------------------
# TEST LIMIT
# ------------------------------------------------------------
#
# Start with 50.
# After validating the results, set this to None to process all mail.
#
PROCESS_LIMIT = 50

# Maximum body characters sent to Qwen.
MAX_BODY_FOR_AI = 5000

# Pause slightly between AI calls.
AI_CALL_DELAY_SECONDS = 0.10

# If True, messages already carrying one of our AI category labels
# will not be sent to Qwen again.
SKIP_ALREADY_CLASSIFIED = True


# ============================================================
# CONTROLLED CATEGORIES
# ============================================================

CATEGORIES = [
    "Business",
    "Finance",
    "Personal",
    "Marketing",
    "Newsletter",
    "OTP",
    "Bills",
    "Orders",
    "Banking",
    "Tax",
    "Government",
    "Legal",
    "Jobs",
    "Social Security",
    "Subscriptions",
    "Notifications",
    "Receipts",
    "Support",
    "Spam",
    "Other",
]


# ============================================================
# GMAIL PERMISSIONS
# ============================================================

# MODIFY allows:
# - reading mail
# - creating/applying labels
# - removing INBOX label (archive)
#
# It does NOT automatically delete anything.
SCOPES = [
    "https://www.googleapis.com/auth/gmail.modify"
]


# ============================================================
# AUTHENTICATION
# ============================================================

def authenticate_gmail():
    creds = None

    if os.path.exists(TOKEN_FILE):
        try:
            creds = Credentials.from_authorized_user_file(
                TOKEN_FILE,
                SCOPES
            )
        except Exception:
            creds = None

    if not creds or not creds.valid:

        if creds and creds.expired and creds.refresh_token:
            print("Refreshing Gmail organizer authentication...")
            creds.refresh(Request())

        else:
            print("\nOrganizer authorization required.")
            print(
                "A browser window will open because this organizer needs "
                "Gmail MODIFY permission."
            )

            if not os.path.exists(CREDENTIALS_FILE):
                raise FileNotFoundError(
                    f"\nGoogle OAuth credentials not found:\n"
                    f"{CREDENTIALS_FILE}\n"
                )

            flow = InstalledAppFlow.from_client_secrets_file(
                CREDENTIALS_FILE,
                SCOPES
            )

            creds = flow.run_local_server(port=0)

        with open(TOKEN_FILE, "w", encoding="utf-8") as token:
            token.write(creds.to_json())

    return build(
        "gmail",
        "v1",
        credentials=creds
    )


# ============================================================
# BASIC HELPERS
# ============================================================

def get_header(headers, name):
    target = name.lower()

    for header in headers:
        if header.get("name", "").lower() == target:
            return header.get("value", "")

    return ""


def parse_email_date(date_string):
    if not date_string:
        return None

    try:
        dt = parsedate_to_datetime(date_string)

        if dt.tzinfo is None:
            dt = dt.replace(tzinfo=timezone.utc)

        return dt

    except Exception:
        return None


def clean_html(html):
    if not html:
        return ""

    html = re.sub(
        r"<script.*?>.*?</script>",
        " ",
        html,
        flags=re.IGNORECASE | re.DOTALL
    )

    html = re.sub(
        r"<style.*?>.*?</style>",
        " ",
        html,
        flags=re.IGNORECASE | re.DOTALL
    )

    html = re.sub(
        r"<br\s*/?>",
        "\n",
        html,
        flags=re.IGNORECASE
    )

    html = re.sub(
        r"</p>",
        "\n",
        html,
        flags=re.IGNORECASE
    )

    html = re.sub(r"<[^>]+>", " ", html)
    html = unescape(html)

    return re.sub(r"\s+", " ", html).strip()


def extract_text_from_payload(payload):
    plain_parts = []
    html_parts = []

    def process_part(part):
        mime_type = part.get("mimeType", "")
        body = part.get("body", {})
        data = body.get("data")

        if data:
            try:
                decoded = base64.urlsafe_b64decode(data).decode(
                    "utf-8",
                    errors="ignore"
                )

                if mime_type == "text/plain":
                    plain_parts.append(decoded)

                elif mime_type == "text/html":
                    html_parts.append(decoded)

            except Exception:
                pass

        for child in part.get("parts", []):
            process_part(child)

    process_part(payload)

    if plain_parts:
        text = "\n".join(plain_parts)

    elif html_parts:
        text = clean_html("\n".join(html_parts))

    else:
        text = ""

    return re.sub(r"\s+", " ", text).strip()


def calculate_age_days(email_date):
    if not email_date:
        return None

    now = datetime.now(timezone.utc)

    try:
        return max(0, (now - email_date.astimezone(timezone.utc)).days)
    except Exception:
        return None


# ============================================================
# GMAIL LABEL MANAGEMENT
# ============================================================

def get_label_maps(service):
    response = service.users().labels().list(
        userId="me"
    ).execute()

    labels = response.get("labels", [])

    name_to_id = {}
    id_to_name = {}

    for label in labels:
        name = label.get("name", "")
        label_id = label.get("id", "")

        if name and label_id:
            name_to_id[name] = label_id
            id_to_name[label_id] = name

    return name_to_id, id_to_name


def create_label(service, label_name):
    result = service.users().labels().create(
        userId="me",
        body={
            "name": label_name,
            "labelListVisibility": "labelShow",
            "messageListVisibility": "show",
        }
    ).execute()

    return result["id"]


def ensure_category_labels(service):
    """
    Creates the parent/category labels only when DRY_RUN is False.
    Returns mapping: category -> Gmail label ID.
    """

    name_to_id, _ = get_label_maps(service)
    category_label_ids = {}

    for category in CATEGORIES:
        full_name = f"{LABEL_ROOT}/{category}"

        if full_name in name_to_id:
            category_label_ids[category] = name_to_id[full_name]
            continue

        if DRY_RUN:
            category_label_ids[category] = None
            print(f"[DRY RUN] Would create Gmail label: {full_name}")
        else:
            print(f"Creating Gmail label: {full_name}")
            label_id = create_label(
                service,
                full_name
            )

            category_label_ids[category] = label_id

    return category_label_ids


def category_from_existing_labels(label_ids, id_to_name):
    """
    If a message already has AI Mail/<Category>, return that category.
    """

    prefix = f"{LABEL_ROOT}/"

    for label_id in label_ids:
        label_name = id_to_name.get(label_id, "")

        if not label_name.startswith(prefix):
            continue

        category = label_name[len(prefix):]

        if category in CATEGORIES:
            return category

    return None


# ============================================================
# MESSAGE LIST
# ============================================================

def get_all_message_ids(service):
    """
    Gets Gmail messages excluding Trash, Spam and Drafts.
    """

    print("\nScanning Gmail...")

    message_ids = []
    page_token = None

    # Exclude system locations we do not want to reorganize.
    query = "-in:trash -in:spam -in:drafts"

    while True:
        response = service.users().messages().list(
            userId="me",
            q=query,
            maxResults=500,
            pageToken=page_token
        ).execute()

        for item in response.get("messages", []):
            message_ids.append(item["id"])

            if (
                PROCESS_LIMIT is not None
                and len(message_ids) >= PROCESS_LIMIT
            ):
                return message_ids[:PROCESS_LIMIT]

        page_token = response.get("nextPageToken")

        if not page_token:
            break

    return message_ids


# ============================================================
# QWEN CATEGORY CLASSIFICATION
# ============================================================

def classify_category_with_qwen(
    sender,
    recipient,
    subject,
    email_date,
    body,
):
    prompt = f"""
You are a Gmail email classification agent.

Classify the email into EXACTLY ONE category.

Return ONLY valid JSON.
Do not use markdown.
Do not include any explanation outside JSON.

Required structure:

{{
  "category": "",
  "reason": "",
  "confidence": 0.0
}}

CATEGORY must be exactly one of:

Business
Finance
Personal
Marketing
Newsletter
OTP
Bills
Orders
Banking
Tax
Government
Legal
Jobs
Social Security
Subscriptions
Notifications
Receipts
Support
Spam
Other

CATEGORY GUIDANCE:

Business:
Client, vendor, professional, company, commercial or work correspondence.

Finance:
General financial matters that are not specifically Banking, Tax,
Bills, Orders or Receipts.

Personal:
Family, friends and genuinely personal correspondence.

Marketing:
Promotions, advertisements, sales, offers and discounts.

Newsletter:
Recurring publications, updates, digests and newsletters.

OTP:
One-time passwords, login verification codes and authentication codes.

Bills:
Bills, invoices and payment-due notices.

Orders:
Order confirmations, shipping, delivery, return and e-commerce messages.

Banking:
Bank accounts, cards, transactions, loans, statements and banking alerts.

Tax:
Income tax, GST, filing, tax notices, tax documents or tax payments.

Government:
Government departments, authorities and official public-service messages.

Legal:
Legal notices, contracts, legal correspondence and legal obligations.

Jobs:
Recruiters, job portals, applications, interviews and job opportunities.

Social Security:
Provident fund, pension, retirement/social-security and official benefit messages.
Do NOT use this for social media.

Subscriptions:
Subscription renewals, cancellation, membership or service subscription notices.

Notifications:
General automated information/alerts that do not fit another category.

Receipts:
Receipts and confirmations of payments already completed.

Support:
Customer support, help desk, complaint, ticket or service-resolution email.

Spam:
Clearly unwanted, deceptive, junk or irrelevant email.

Other:
Only if no other category fits.

Important:
- Pick the BEST single category.
- Do not invent categories.
- Confidence must be a number from 0 to 1.

EMAIL DATE:
{email_date.strftime("%Y-%m-%d %H:%M:%S") if email_date else "Unknown"}

FROM:
{sender}

TO:
{recipient}

SUBJECT:
{subject}

BODY:
{body[:MAX_BODY_FOR_AI]}
"""

    try:
        response = requests.post(
            OLLAMA_URL,
            json={
                "model": OLLAMA_MODEL,
                "prompt": prompt,
                "stream": False,
                "format": "json",
            },
            timeout=180
        )

        response.raise_for_status()

        raw_response = response.json().get("response", "")

        if not raw_response:
            raise ValueError("Empty Qwen response")

        result = json.loads(raw_response)

        raw_category = str(
            result.get("category", "")
        ).strip()

        category = "Other"

        for allowed in CATEGORIES:
            if raw_category.lower() == allowed.lower():
                category = allowed
                break

        try:
            confidence = float(
                result.get("confidence", 0)
            )
        except Exception:
            confidence = 0.0

        confidence = max(
            0.0,
            min(confidence, 1.0)
        )

        reason = str(
            result.get("reason", "")
        ).strip()

        return {
            "category": category,
            "confidence": confidence,
            "reason": reason,
        }

    except Exception as error:
        return {
            "category": "Other",
            "confidence": 0.0,
            "reason": f"AI classification error: {error}",
        }


# ============================================================
# MODIFY GMAIL
# ============================================================

def apply_category_label(
    service,
    message_id,
    target_category,
    category_label_ids,
    existing_category,
):
    """
    Applies exactly one AI Mail category label.
    If another AI category is already present, removes it.
    """

    target_label_id = category_label_ids.get(
        target_category
    )

    if DRY_RUN:
        print(
            f"[DRY RUN] Would label as: "
            f"{LABEL_ROOT}/{target_category}"
        )
        return

    if not target_label_id:
        raise RuntimeError(
            f"Missing Gmail label ID for category: {target_category}"
        )

    # Refresh label map because we need category label IDs
    # to remove any old AI category cleanly.
    name_to_id, _ = get_label_maps(service)

    remove_ids = []

    for category in CATEGORIES:
        if category == target_category:
            continue

        label_name = f"{LABEL_ROOT}/{category}"
        old_id = name_to_id.get(label_name)

        if old_id:
            remove_ids.append(old_id)

    service.users().messages().modify(
        userId="me",
        id=message_id,
        body={
            "addLabelIds": [target_label_id],
            "removeLabelIds": remove_ids,
        }
    ).execute()


def archive_if_old(
    service,
    message_id,
    gmail_label_ids,
    age_days,
):
    """
    Removes INBOX label only when:
    - message currently has INBOX
    - email age is greater than retention days
    """

    if age_days is None:
        print("Archive rule: skipped because email date is unknown.")
        return False

    in_inbox = "INBOX" in gmail_label_ids

    if not in_inbox:
        print("Archive rule: message is already outside Inbox.")
        return False

    if age_days <= INBOX_RETENTION_DAYS:
        print(
            f"Archive rule: keep in Inbox "
            f"({age_days} days old)."
        )
        return False

    if DRY_RUN:
        print(
            f"[DRY RUN] Would archive: "
            f"{age_days} days old (> {INBOX_RETENTION_DAYS})."
        )
        return True

    service.users().messages().modify(
        userId="me",
        id=message_id,
        body={
            "removeLabelIds": ["INBOX"]
        }
    ).execute()

    print(
        f"Archived from Inbox: "
        f"{age_days} days old."
    )

    return True


# ============================================================
# PROCESS ONE MESSAGE
# ============================================================

def process_message(
    service,
    message_id,
    category_label_ids,
    id_to_name,
):
    message = service.users().messages().get(
        userId="me",
        id=message_id,
        format="full"
    ).execute()

    payload = message.get("payload", {})
    headers = payload.get("headers", [])
    gmail_label_ids = message.get("labelIds", [])

    sender = get_header(headers, "From")
    recipient = get_header(headers, "To")
    subject = get_header(headers, "Subject")
    date_string = get_header(headers, "Date")

    email_date = parse_email_date(date_string)
    age_days = calculate_age_days(email_date)
    body = extract_text_from_payload(payload)

    existing_category = category_from_existing_labels(
        gmail_label_ids,
        id_to_name
    )

    print("\n" + "-" * 72)
    print(f"Subject : {subject}")
    print(f"Sender  : {sender}")
    print(
        f"Age     : "
        f"{age_days if age_days is not None else 'Unknown'} days"
    )

    if (
        SKIP_ALREADY_CLASSIFIED
        and existing_category
    ):
        category = existing_category
        confidence = 1.0
        reason = "Already classified by this organizer."

        print(
            f"Category: {category} "
            f"(existing AI label)"
        )

    else:
        print("Sending to Qwen for category classification...")

        ai = classify_category_with_qwen(
            sender=sender,
            recipient=recipient,
            subject=subject,
            email_date=email_date,
            body=body,
        )

        category = ai["category"]
        confidence = ai["confidence"]
        reason = ai["reason"]

        print(f"Category   : {category}")
        print(f"Confidence : {confidence:.2f}")
        print(f"Reason     : {reason}")

        apply_category_label(
            service=service,
            message_id=message_id,
            target_category=category,
            category_label_ids=category_label_ids,
            existing_category=existing_category,
        )

    archived = archive_if_old(
        service=service,
        message_id=message_id,
        gmail_label_ids=gmail_label_ids,
        age_days=age_days,
    )

    inbox_status = (
        "In Inbox"
        if "INBOX" in gmail_label_ids
        else "Outside Inbox"
    )

    if age_days is None:
        archive_decision = "Unknown Date"
    elif "INBOX" not in gmail_label_ids:
        archive_decision = "Already Outside Inbox"
    elif age_days > INBOX_RETENTION_DAYS:
        archive_decision = "Archive"
    else:
        archive_decision = "Keep in Inbox"

    return {
        "message_id": message_id,
        "date": (
            email_date.strftime("%Y-%m-%d %H:%M:%S")
            if email_date else ""
        ),
        "subject": subject,
        "sender": sender,
        "recipient": recipient,
        "age_days": age_days,
        "category": category,
        "confidence": confidence,
        "reason": reason,
        "existing_category": existing_category or "",
        "inbox_status": inbox_status,
        "archive_decision": archive_decision,
        "archived": archived,
        "already_classified": bool(existing_category),
    }



# ============================================================
# EXCEL REVIEW REPORT
# ============================================================

def create_excel_report(records):
    timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
    mode = "DRY_RUN" if DRY_RUN else "LIVE"
    filepath = os.path.join(
        REPORT_DIR,
        f"Email_Organizer_V2_{mode}_{timestamp}.xlsx"
    )

    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"

    summary.append(["SRNR GMAIL AI ORGANIZER - V2"])
    summary.append([])
    summary.append([
        "Run Mode",
        "DRY RUN - NO GMAIL CHANGES"
        if DRY_RUN else
        "LIVE - GMAIL CHANGES ENABLED"
    ])
    summary.append(["Messages Processed", len(records)])
    summary.append([
        "Archive Candidates",
        sum(1 for r in records if r["archive_decision"] == "Archive")
    ])
    summary.append([
        "Keep In Inbox",
        sum(1 for r in records if r["archive_decision"] == "Keep in Inbox")
    ])
    summary.append([])
    summary.append(["CATEGORY", "COUNT"])

    counts = {}
    for r in records:
        counts[r["category"]] = counts.get(r["category"], 0) + 1

    for category in CATEGORIES:
        summary.append([category, counts.get(category, 0)])

    details = wb.create_sheet("Email Review")
    headers = [
        "Message ID",
        "Date",
        "Age (Days)",
        "Sender",
        "Recipient",
        "Subject",
        "Category",
        "Confidence",
        "AI Reason",
        "Existing AI Category",
        "Inbox Status",
        "Archive Decision",
    ]
    details.append(headers)

    archive_sheet = wb.create_sheet("Archive Candidates")
    archive_sheet.append(headers)

    for r in records:
        row = [
            r["message_id"],
            r["date"],
            r["age_days"],
            r["sender"],
            r["recipient"],
            r["subject"],
            r["category"],
            r["confidence"],
            r["reason"],
            r["existing_category"],
            r["inbox_status"],
            r["archive_decision"],
        ]
        details.append(row)

        if r["archive_decision"] == "Archive":
            archive_sheet.append(row)

    for sheet in wb.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

        for cell in sheet[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="D9EAF7"
            )

        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

        for column in sheet.columns:
            letter = get_column_letter(column[0].column)
            max_len = 0

            for cell in column:
                try:
                    max_len = max(max_len, len(str(cell.value)))
                except Exception:
                    pass

            sheet.column_dimensions[letter].width = min(
                max(max_len + 2, 12),
                55
            )

    wb.save(filepath)
    return filepath


# ============================================================
# MAIN
# ============================================================

def main():
    print("=" * 72)
    print("SRNR GMAIL AI ORGANIZER - VERSION 2")
    print("=" * 72)

    print(f"\nProject directory : {PROJECT_DIR}")
    print(f"AI model          : {OLLAMA_MODEL}")
    print(f"Gmail label root  : {LABEL_ROOT}")
    print(f"Inbox retention   : {INBOX_RETENTION_DAYS} days")
    print(f"Dry run           : {DRY_RUN}")
    print(
        f"Process limit     : "
        f"{PROCESS_LIMIT if PROCESS_LIMIT is not None else 'ALL'}"
    )

    print("\nSAFETY:")
    print("  - Emails are NEVER deleted.")
    print("  - Trash is NOT processed.")
    print("  - Spam is NOT processed.")
    print("  - Drafts are NOT processed.")
    print("  - Older messages are only archived by removing INBOX.")
    print("  - Category labels are applied under the AI Mail parent.")

    if DRY_RUN:
        print("\n*** DRY RUN MODE ***")
        print("Gmail will NOT be changed during this run.")

    # --------------------------------------------------------
    # OLLAMA CHECK
    # --------------------------------------------------------

    print("\nChecking Ollama...")

    try:
        ollama_test = requests.get(
            "http://127.0.0.1:11434/api/tags",
            timeout=10
        )

        ollama_test.raise_for_status()

        installed_models = [
            model.get("name", "")
            for model in ollama_test.json().get(
                "models",
                []
            )
        ]

        print("Ollama connection: OK")

        if OLLAMA_MODEL not in installed_models:
            print(
                f"\nERROR: Configured model "
                f"'{OLLAMA_MODEL}' is not installed."
            )

            print("Installed models:")

            for model in installed_models:
                print(f"  - {model}")

            return

    except Exception as error:
        print("\nERROR: Cannot connect to Ollama.")
        print(f"Details: {error}")
        return

    # --------------------------------------------------------
    # GMAIL
    # --------------------------------------------------------

    print("\nConnecting to Gmail...")

    service = authenticate_gmail()

    print("Gmail connection: OK")

    # --------------------------------------------------------
    # LABELS
    # --------------------------------------------------------

    category_label_ids = ensure_category_labels(
        service
    )

    _, id_to_name = get_label_maps(
        service
    )

    # --------------------------------------------------------
    # MESSAGES
    # --------------------------------------------------------

    message_ids = get_all_message_ids(
        service
    )

    print(
        f"\nMessages selected: "
        f"{len(message_ids):,}"
    )

    if not message_ids:
        print("No messages found.")
        return

    # --------------------------------------------------------
    # PROCESS
    # --------------------------------------------------------

    records = []

    for index, message_id in enumerate(
        message_ids,
        start=1
    ):
        print(
            f"\nProcessing email "
            f"{index}/{len(message_ids)}"
        )

        try:
            record = process_message(
                service=service,
                message_id=message_id,
                category_label_ids=category_label_ids,
                id_to_name=id_to_name,
            )

            records.append(record)

        except Exception as error:
            print(
                f"ERROR processing message "
                f"{message_id}: {error}"
            )

        if AI_CALL_DELAY_SECONDS:
            time.sleep(
                AI_CALL_DELAY_SECONDS
            )

    # --------------------------------------------------------
    # EXCEL REPORT
    # --------------------------------------------------------

    print("\nCreating Excel review report...")
    report_path = create_excel_report(records)
    print(f"Excel report: {report_path}")

    # --------------------------------------------------------
    # SUMMARY
    # --------------------------------------------------------

    archived_count = sum(
        1 for r in records
        if r["archived"]
    )

    already_classified_count = sum(
        1 for r in records
        if r["already_classified"]
    )

    category_counts = {}

    for record in records:
        category_counts[record["category"]] = (
            category_counts.get(
                record["category"],
                0
            ) + 1
        )

    print("\n" + "=" * 72)
    print("ORGANIZER RUN COMPLETE")
    print("=" * 72)

    print(
        f"\nMessages processed      : "
        f"{len(records):,}"
    )

    print(
        f"Already AI-classified   : "
        f"{already_classified_count:,}"
    )

    if DRY_RUN:
        print(
            f"Would archive from Inbox: "
            f"{archived_count:,}"
        )
    else:
        print(
            f"Archived from Inbox     : "
            f"{archived_count:,}"
        )

    print("\nCategory totals:")

    for category in CATEGORIES:
        count = category_counts.get(
            category,
            0
        )

        if count:
            print(
                f"  {category:<18} {count}"
            )

    if DRY_RUN:
        print("\nNo Gmail changes were made.")
        print(
            "After checking these results, change "
            "DRY_RUN = False to enable organization."
        )
    else:
        print("\nGmail organization changes were applied.")
        print("No emails were deleted.")


# ============================================================
# START
# ============================================================

if __name__ == "__main__":
    main()
